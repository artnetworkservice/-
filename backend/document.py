"""Create an editable DOCX from the supplied 51-page source document."""

from __future__ import annotations

from copy import deepcopy
from datetime import date
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
import re
import shutil
import subprocess
import tempfile
import os
import sys
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from .pricing import (
    calculate_document,
    compact_item_references,
    format_money,
    format_quantity,
    thai_baht_words,
    thai_number,
)


PROJECT_DIR = Path(getattr(sys, "_MEIPASS", Path(__file__).resolve().parents[1]))
DEFAULT_SOURCE = next(
    path for path in PROJECT_DIR.glob("*.docx")
    if path.name.startswith("ราคากลางงานจัดซื้อพัสดุ ") and not path.name.startswith("~$")
)
THAI_MONTHS = ("ม.ค.", "ก.พ.", "มี.ค.", "เม.ย.", "พ.ค.", "มิ.ย.", "ก.ค.", "ส.ค.", "ก.ย.", "ต.ค.", "พ.ย.", "ธ.ค.")
SUBJECT_PREFIX = "ราคากลางจัดซื้อพัสดุ"
ENGINE_REVISION = "1.1.3"


def thai_date(raw: str) -> str:
    if not raw:
        return ""
    try:
        value = date.fromisoformat(raw)
    except ValueError:
        return raw
    return f"{thai_number(value.day)} {THAI_MONTHS[value.month - 1]}{thai_number(f'{(value.year + 543) % 100:02d}')}"


def _set_text(paragraph, text: str, *, preserve_drawings: bool = False) -> None:
    """Keep paragraph geometry; optionally keep graphic runs used in the header."""
    element = paragraph._p
    first_rpr = None
    for run in paragraph.runs:
        if first_rpr is None and run._r.rPr is not None and run.text.strip():
            first_rpr = deepcopy(run._r.rPr)
    for child in list(element):
        if child.tag == qn("w:pPr"):
            continue
        if preserve_drawings and any(node.tag == qn("w:drawing") for node in child.iter()):
            continue
        element.remove(child)
    run = OxmlElement("w:r")
    if first_rpr is not None:
        run.append(first_rpr)
    parts = re.split(r"(\n|\t)", text)
    for part in parts:
        if part == "\n":
            run.append(OxmlElement("w:br"))
        elif part == "\t":
            run.append(OxmlElement("w:tab"))
        elif part:
            node = OxmlElement("w:t")
            node.set(qn("xml:space"), "preserve")
            node.text = part
            run.append(node)
    element.append(run)


def _new_paragraph(exemplar, text: str, *, continuation: bool = False):
    copied = deepcopy(exemplar._p)
    from docx.text.paragraph import Paragraph
    wrapper = Paragraph(copied, exemplar._parent)
    _set_text(wrapper, text)
    wrapper.alignment = WD_ALIGN_PARAGRAPH.THAI_JUSTIFY
    if continuation:
        wrapper.paragraph_format.first_line_indent = 0
        wrapper.paragraph_format.space_before = 0
        wrapper.paragraph_format.space_after = 0
    return copied


def _replace_header_text(paragraph, text: str) -> None:
    _set_text(paragraph, text, preserve_drawings=True)


def _set_header_run(paragraph, index: int, value: str, *, clear_after: bool = True) -> None:
    runs = paragraph.runs
    runs[index].text = value
    if clear_after:
        for run in runs[index + 1:]:
            if not any(node.tag == qn("w:drawing") for node in run._r.iter()):
                run.text = ""


def _subject_detail(subject: str) -> str:
    return re.sub(r"^ราคากลาง(?:งาน)?จัดซื้อ(?:พัสดุ)?\s*", "", subject).strip()


def _procurement_subject(subject: str) -> str:
    detail = _subject_detail(subject)
    return f"พัสดุ {detail}".strip() if subject.startswith(SUBJECT_PREFIX) else detail


def _render_template(text: str, values: dict[str, str]) -> str:
    for key, value in values.items():
        text = text.replace("{{" + key + "}}", value)
    return text


def _default_blocks(source: Document) -> dict[str, str]:
    p = source.paragraphs
    if len(p) < 100:
        if len(p) < 18 or not p[7].text.strip().startswith("๒.๑") or not p[12].text.strip().startswith("๓."):
            raise ValueError("แม่แบบ Word ต้องมีหัวเอกสาร รายการตัวอย่าง ข้อสรุป และช่องลงชื่อในโครงสร้างที่รองรับ")
        standard = _default_blocks(Document(DEFAULT_SOURCE))
        return {
            **standard,
            "agency": p[1].text.replace("ส่วนราชการ", "", 1).strip(),
            "intro": p[5].text,
            "no_prior_clause": p[7].text.split("ตามเรื่องนี้", 1)[-1].strip(),
            "closing": p[13].text,
            "signature_1": p[15].text.strip(),
            "signature_2": p[16].text.strip(),
            "signature_3": p[17].text.strip(),
        }
    return {
        "agency": p[1].text.replace("ส่วนราชการ", "").strip(),
        "intro": p[5].text,
        "prior_clause": p[7].text.split("ตามเรื่องนี้", 1)[1].strip(),
        "no_prior_clause": p[607].text.split("ตามเรื่องนี้", 1)[1].strip(),
        "reference_announcement": "ตามบัญชีแสดงราคาอ้างอิงและรายละเอียดคุณลักษณะเฉพาะของพัสดุที่กำหนดราคาอ้างอิงท้ายประกาศกรมบัญชีกลาง เรื่องกำหนดราคาอ้างอิงพัสดุ พ.ศ.๒๕๖๖ ประกาศ ณ วันที่ ๑๒ กรกฎาคม พ.ศ.๒๕๖๖",
        "closing": p[680].text,
        "signature_1": p[682].text.strip(),
        "signature_2": p[683].text.strip(),
        "signature_3": p[684].text.strip(),
    }


def default_blocks() -> dict[str, str]:
    return _default_blocks(Document(DEFAULT_SOURCE))


def _item_text(item: dict[str, Any]) -> dict[str, str]:
    qty = format_quantity(item["quantity"])
    unit = item["unit"]
    price = format_money(item["selected_unit_price"])
    total = format_money(item["total"])
    return {"name": item["name"], "quantity": qty, "unit": unit, "price": price, "total": total}


def _summary(items: list[dict[str, Any]], subject_detail: str, grand_total, blocks: dict[str, str]) -> str:
    equal = [x["index"] for x in items if x["decision"] == "equal"]
    reference = [x["index"] for x in items if x["decision"] == "reference"]
    market = [x["index"] for x in items if x["decision"] in ("market_only", "market_lower")]
    clauses = []
    if equal:
        clauses.append(
            "ราคาที่เคยจัดซื้อครั้งหลังสุดภายในระยะเวลา ๒ ปีงบประมาณ ตามข้อ "
            + compact_item_references(equal, suffix=".1")
            + " เท่ากับราคาที่ได้จากการสืบราคา ฯ ตามข้อ "
            + compact_item_references(equal, suffix=".2")
        )
    if reference:
        clauses.append(
            "ตามข้อ " + compact_item_references(reference)
            + " ใช้ราคา" + " ".join(blocks["reference_announcement"].split())
        )
    if market:
        references = compact_item_references(market)
        quote_counts = [len(x["quotes"]) for x in items if x["index"] in market]
        count_text = thai_number(quote_counts[0]) if len(set(quote_counts)) == 1 else "ตามแต่ละรายการ"
        clauses.append(
            f"ผู้ยื่นเสนอราคาทั้ง {count_text} ราย เสนอราคาถูกต้องมีรายละเอียดพัสดุครบถ้วนตามขอบเขตงาน "
            f"ตามข้อ {references} เห็นสมควรใช้ราคาต่ำสุดที่ได้จากการสืบราคา ฯ"
        )
    total = format_money(grand_total)
    words = thai_baht_words(grand_total)
    return (
        "๓.  คณก.ฯ พิจารณาตามข้อ ๒ แล้ว "
        + " และ ".join(clauses)
        + f" ในการจัดซื้อพัสดุ {subject_detail} รวมเป็นเงินทั้งสิ้น {total} บาท "
        + f"({words})เป็นราคารวมภาษีมูลค่าเพิ่ม ตลอดจนภาษีอากรอื่นๆ และค่าใช้จ่ายทั้งปวงด้วยแล้ว"
        + "ซึ่งเป็นตาม พ.ร.บ.ฯ มาตรา ๔ (๔) เป็นราคากลางในการจัดซื้อ"
        " "
        " "
    )


def create_docx(data: dict[str, Any], destination: Path, *, source_path: Path | None = None, blocks: dict[str, str] | None = None) -> dict[str, Any]:
    result = calculate_document(data)
    source = Document(source_path or DEFAULT_SOURCE)
    original = list(source.paragraphs)
    compact = len(original) < 100
    if compact and (len(original) < 18 or not original[12].text.strip().startswith("๓.")):
        raise ValueError("แม่แบบ Word ไม่ตรงกับโครงสร้างที่รองรับ")
    if not compact and len(original) < 685:
        raise ValueError("แม่แบบ Word ไม่ตรงกับโครงสร้างที่รองรับ")
    settings = _default_blocks(source)
    settings.update(blocks or {})
    body = source._element.body
    summary_index = 12 if compact else 679
    anchor = original[summary_index]._p
    samples = ({index: original[7 if index in (7, 57, 607) else 11 if index in (16, 58, 611) else 8]
                for index in (7, 8, 9, 10, 16, 57, 58, 607, 608, 611)} if compact else
               {index: original[index] for index in (7, 8, 9, 10, 16, 57, 58, 607, 608, 611)})

    # Remove the 70 sample items and editorial placeholders, then insert new blocks.
    for paragraph in original[7:summary_index]:
        body.remove(paragraph._p)

    if compact:
        if settings["agency"] != original[1].text.replace("ส่วนราชการ", "", 1).strip():
            _set_header_run(original[1], 3, settings["agency"])
        original[2].runs[3].text = "\t" + thai_number(result["document_number"])
        original[2].runs[4].text = "\t"
        original[2].runs[6].text = "\t" + thai_date(result["date"])
        for run in original[2].runs[7:]:
            run.text = ""
        original[3].runs[2].text = "\t" + result["subject"]
    else:
        if settings["agency"] != original[1].text.replace("ส่วนราชการ", "").strip():
            _set_header_run(original[1], 3, settings["agency"])
        original[2].runs[3].text = "\t" + thai_number(result["document_number"])
        original[2].runs[4].text = "\t"
        original[2].runs[6].text = "    " + thai_date(result["date"])
        for run in original[2].runs[7:]:
            run.text = ""
        _set_header_run(original[3], 4, result["subject"])
    if settings["intro"] != original[5].text:
        _set_text(original[5], settings["intro"])
    procurement = _procurement_subject(result["subject"])
    _set_text(original[6], f"๒.  คณก.ฯ ได้ตรวจสอบและดำเนินการแล้ว ในการจัดซื้อ{procurement} จำนวน {thai_number(len(result['items']))} รายการ ดังนี้")
    original[6].alignment = WD_ALIGN_PARAGRAPH.THAI_JUSTIFY

    def add(text: str, sample_index: int) -> None:
        for position, line in enumerate(text.split("\n")):
            anchor.addprevious(_new_paragraph(samples[sample_index], line, continuation=position > 0))

    for item in result["items"]:
        n = thai_number(item["index"])
        basics = _item_text(item)
        name = basics["name"]
        qty = basics["quantity"]
        unit = basics["unit"]
        price = basics["price"]
        total = basics["total"]
        if item["source"] == "reference":
            announcement = settings["reference_announcement"]
            add(f"๒.{n}  รายการที่ {n} {name} ใช้ราคา{announcement} ในราคา{unit}ละ {price} บาท จำนวน {qty} {unit} รวมเป็นเงินทั้งสิ้น {total} บาท", 57)
            add(f"คณก.ตรวจสอบแล้ว ในการจัดซื้อ{name} ใช้ราคา{announcement} ในราคา{unit}ละ {price} บาท จำนวน {qty} {unit} รวมเป็นเงินทั้งสิ้น {total} บาท", 58)
            continue

        has_prior = item["source"] == "prior"
        clause = settings["prior_clause"] if has_prior else settings["no_prior_clause"]
        add(f"๒.{n}  รายการที่ {n} {name} จำนวน {qty} {unit} ตามเรื่องนี้\n{clause}", 7 if has_prior else 607)
        if has_prior:
            prior = item["prior"]
            date_part = " ลง " + thai_date(prior["document_date"]) if prior["document_date"] else ""
            old_price = format_money(prior["unit_price"])
            old_total = format_money((prior["unit_price"] * item["quantity"]).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))
            add(
                f"๒.{n}.๑  ราคาที่เคยจัดซื้อครั้งหลังสุดภายในระยะเวลา ๒ ปีงบประมาณ จัดซื้อตาม "
                f"{thai_number(prior['document_number'])}{date_part} ในราคา{unit}ละ {old_price} บาท "
                f"ในการจัดซื้อครั้งนี้ จำนวน {qty} {unit} รวมเป็นเงินทั้งสิ้น {old_total} บาท",
                8,
            )
            add(f"๒.{n}.๒  ได้สืบราคาจากผู้ประกอบการ จำนวน {thai_number(len(item['quotes']))} ราย ประกอบด้วย", 9)
        for quote_index, quote in enumerate(item["quotes"], 1):
            subnumber = f"๒.{n}.๒.{thai_number(quote_index)}" if has_prior else f"๒.{n}.{thai_number(quote_index)}"
            quote_total = format_money((quote["unit_price"] * item["quantity"]).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))
            add(
                f"{subnumber}  {quote['company']} เสนอราคาในการจัดซื้อ {name} ในราคา{unit}ละ "
                f"{format_money(quote['unit_price'])} บาท จำนวน {qty} {unit} รวมเป็นเงินทั้งสิ้น {quote_total} บาท",
                10 if has_prior else 608,
            )
        add(
            f"คณก.ฯ ตรวจสอบแล้ว ผู้ยื่นเสนอราคาทั้ง {thai_number(len(item['quotes']))} ราย เสนอราคาถูกต้อง "
            f"มีรายละเอียดพัสดุครบถ้วนตามขอบเขตงานซึ่งราคาต่ำสุดที่ได้จากการสืบราคา ฯ "
            f"ในการจัดซื้อ{name} ในราคา{unit}ละ {price} บาท จำนวน {qty} {unit} รวมเป็นเงินทั้งสิ้น {total} บาท",
            16 if has_prior else 611,
        )

    _set_text(original[summary_index], _summary(result["items"], _subject_detail(result["subject"]), result["total"], settings))
    original[summary_index].alignment = WD_ALIGN_PARAGRAPH.THAI_JUSTIFY
    closing_index = 13 if compact else 680
    if settings["closing"] != original[closing_index].text:
        _set_text(original[closing_index], settings["closing"])
    signature_indices = (15, 16, 17) if compact else (682, 683, 684)
    for para, field in zip((original[index] for index in signature_indices), ("signature_1", "signature_2", "signature_3")):
        if settings[field] != para.text.strip():
            _set_text(para, settings[field])
    for paragraph in original[signature_indices[-1] + 1:]:
        if not paragraph.text.strip() and not any(
            node.tag in (qn("w:drawing"), qn("w:sectPr")) for node in paragraph._p.iter()
        ):
            body.remove(paragraph._p)
    destination.parent.mkdir(parents=True, exist_ok=True)
    source.save(destination)
    return result


def convert_to_pdf(docx_path: Path, destination: Path) -> None:
    office = shutil.which("soffice") or shutil.which("libreoffice")
    if not office:
        candidates = [
            Path("/Applications/LibreOffice.app/Contents/MacOS/soffice"),
            Path("/Applications/LibreOffice.app/Contents/MacOS/soffice.bin"),
            Path(os.environ.get("PROGRAMFILES", "C:/Program Files")) / "LibreOffice" / "program" / "soffice.exe",
            Path(os.environ.get("PROGRAMFILES(X86)", "C:/Program Files (x86)")) / "LibreOffice" / "program" / "soffice.exe",
        ]
        office = next((str(path) for path in candidates if path.is_file()), None)
    if not office:
        raise RuntimeError("ไม่พบ LibreOffice สำหรับสร้าง PDF กรุณาติดตั้ง LibreOffice")
    destination.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="quotation-office-") as temp:
        profile = Path(temp) / "profile"
        outdir = Path(temp) / "pdf"
        outdir.mkdir()
        environment = os.environ.copy()
        if os.sys.platform == "darwin":
            font_directories = [
                Path("/System/Library/Fonts"),
                Path("/System/Library/Fonts/Supplemental"),
                Path("/Library/Fonts"),
                Path.home() / "Library" / "Fonts",
                Path("/Applications/Microsoft Word.app/Contents/Resources/DFonts"),
                PROJECT_DIR / "frontend" / "public" / "fonts",
            ]
            config = Path(temp) / "fonts.conf"
            dirs = "\n".join(f"<dir>{str(folder).replace('&', '&amp;')}</dir>" for folder in font_directories if folder.exists())
            config.write_text(f'<?xml version="1.0"?><fontconfig>{dirs}<cachedir>{temp}/fontcache</cachedir></fontconfig>', encoding="utf-8")
            environment["FONTCONFIG_FILE"] = str(config)
            (Path(temp) / "fontcache").mkdir()
        command = [office, f"-env:UserInstallation=file://{profile}", "--headless", "--convert-to", "pdf:writer_pdf_Export", "--outdir", str(outdir), str(docx_path)]
        process = subprocess.run(command, capture_output=True, text=True, timeout=120, env=environment)
        generated = outdir / (docx_path.stem + ".pdf")
        if process.returncode != 0 or not generated.exists():
            raise RuntimeError("สร้าง PDF ไม่สำเร็จ: " + (process.stderr or process.stdout).strip()[:500])
        shutil.copy2(generated, destination)
