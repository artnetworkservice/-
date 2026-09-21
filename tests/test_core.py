from decimal import Decimal
from pathlib import Path
import os
import tempfile
import unittest
from zipfile import ZipFile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from backend.document import create_docx
from backend.pricing import ValidationError, calculate_document, compact_item_references, thai_baht_words
from backend.storage import Store


def base_item(source="market"):
    return {
        "name": "ยางลบดินสอ",
        "quantity": "3",
        "unit": "กล่อง",
        "source": source,
        "quotes": [
            {"company": "ร้าน ก", "unit_price": "180"},
            {"company": "ร้าน ข", "unit_price": "185"},
        ],
        "prior": {"document_number": "1841/68", "document_date": "2025-09-11", "quantity": "3", "unit_price": "180"},
        "reference": {"quantity": "3", "unit_price": "102"},
    }


def document(items):
    return {"subject": "ราคากลางงานจัดซื้อวัสดุสำนักงาน", "document_number": "123/69", "date": "2026-09-16", "items": items}


class PricingTests(unittest.TestCase):
    def test_equal_prior_and_quote(self):
        result = calculate_document(document([base_item("prior")]))
        self.assertEqual(result["items"][0]["decision"], "equal")
        self.assertEqual(result["total"], Decimal("540.00"))

    def test_lower_quote_beats_prior(self):
        item = base_item("prior")
        item["prior"]["unit_price"] = "200"
        self.assertEqual(calculate_document(document([item]))["items"][0]["decision"], "market_lower")

    def test_prior_below_every_quote_requires_policy(self):
        item = base_item("prior")
        item["prior"]["unit_price"] = "100"
        with self.assertRaisesRegex(ValidationError, "ต้องยืนยันกติกา"):
            calculate_document(document([item]))

    def test_reference_needs_no_quotes_and_uses_exact_cents(self):
        item = base_item("reference")
        item["quotes"] = []
        item["quantity"] = "2.5"
        item["reference"] = {"quantity": "2.5", "unit_price": "102.20"}
        result = calculate_document(document([item]))
        self.assertEqual(result["items"][0]["decision"], "reference")
        self.assertEqual(result["total"], Decimal("255.50"))

    def test_thai_words_and_refs(self):
        self.assertEqual(thai_baht_words(Decimal("67000")), "หกหมื่นเจ็ดพันบาทถ้วน")
        self.assertEqual(thai_baht_words(Decimal("255.50")), "สองร้อยห้าสิบห้าบาทห้าสิบสตางค์")
        self.assertEqual(compact_item_references([1, 2, 3, 6], suffix=".1"), "๒.๑.๑ - ๒.๓.๑ , ๒.๖.๑")


class DocumentTests(unittest.TestCase):
    def test_generated_docx_replaces_sample_items_and_keeps_garuda(self):
        with tempfile.TemporaryDirectory() as temporary:
            target = Path(temporary) / "test.docx"
            create_docx(document([base_item("prior")]), target)
            generated = Document(target)
            content = "\n".join(p.text for p in generated.paragraphs)
            self.assertIn("จำนวน ๑ รายการ", content)
            self.assertIn("รวมเป็นเงินทั้งสิ้น ๕๔๐ บาท", content)
            self.assertNotIn("รายการที่ ๗๐", content)
            self.assertIn("น.อ.หญิง", content)
            self.assertEqual(next(p for p in generated.paragraphs if "รายการที่ ๑" in p.text).alignment, WD_ALIGN_PARAGRAPH.THAI_JUSTIFY)
            self.assertIn("กรรมการ", generated.paragraphs[-1].text)
            with ZipFile(target) as archive:
                self.assertTrue(any(name.startswith("word/media/") for name in archive.namelist()))

    def test_template_versions_do_not_change_existing_document(self):
        with tempfile.TemporaryDirectory() as temporary:
            store = Store(Path(temporary))
            first = store.list_templates()[0]
            saved = store.save_document({"template_id": first["id"], **document([base_item()])})
            updated = store.update_template(first["id"], "แก้ไขชื่อ", {"agency": "หน่วยงานใหม่"})
            self.assertNotEqual(saved["template_version_id"], updated["id"])
            copied = store.duplicate_document(saved["id"])
            self.assertEqual(copied["template_version_id"], saved["template_version_id"])
            store.set_deleted(saved["id"], True)
            self.assertEqual(len(store.list_documents()), 1)
            self.assertEqual(len(store.list_documents(include_deleted=True)), 1)

    def test_backup_restores_template_and_document_on_another_data_path(self):
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            original = Store(root / "first")
            template = original.list_templates()[0]
            saved = original.save_document({"template_id": template["id"], **document([base_item()])})
            old_environment = os.environ.get("QUOTATION_DATA_DIR")
            os.environ["QUOTATION_DATA_DIR"] = str(root / "test-server")
            try:
                from backend import server
                old_store = server.STORE
                try:
                    server.STORE = original
                    backup = server.backup_bytes()
                    moved = Store(root / "second")
                    server.STORE = moved
                    server.restore_backup(backup)
                    restored = moved.get_document(saved["id"])
                    self.assertEqual(restored["subject"], saved["subject"])
                    version = moved.template_version(restored["template_version_id"])
                    self.assertTrue(Path(version["source_path"]).is_file())
                    self.assertEqual(len(moved.list_documents()), 1)
                finally:
                    server.STORE = old_store
            finally:
                if old_environment is None:
                    os.environ.pop("QUOTATION_DATA_DIR", None)
                else:
                    os.environ["QUOTATION_DATA_DIR"] = old_environment

    def test_import_both_v11_templates_export_and_soft_delete(self):
        sources = ("ต้นฉบับ สายช่างโยธา 70.docx", "ต้นฉบับ สายพลาธิการ 69.docx")
        with tempfile.TemporaryDirectory() as temporary:
            store = Store(Path(temporary))
            for source_name in sources:
                imported = store.import_template(source_name.removesuffix(".docx"), Path(source_name).read_bytes())
                data = {"template_id": imported["template_id"], **document([base_item("market"), base_item("prior"), base_item("reference")])}
                data["items"][1]["prior"]["document_number"] = "เอกสาร 100/68 ลง 11 ก.ย. 68"
                data["items"][1]["prior"]["document_date"] = ""
                saved = store.save_document(data)
                target = Path(temporary) / f"{imported['id']}-result.docx"
                create_docx(data, target, source_path=Path(imported["source_path"]), blocks=imported["blocks"])
                output = Document(target)
                text = "\n".join(p.text for p in output.paragraphs)
                self.assertIn(imported["blocks"]["agency"].split()[0], text)
                self.assertIn("เอกสาร ๑๐๐/๖๘ ลง ๑๑ ก.ย. ๖๘", text)
                self.assertNotIn("ทอ.๗๔ ที่", text)
                self.assertEqual(next(p for p in output.paragraphs if "รายการที่ ๑" in p.text).alignment, WD_ALIGN_PARAGRAPH.THAI_JUSTIFY)
                self.assertTrue(any(row["template_name"] == source_name.removesuffix(".docx") for row in store.list_documents()))
                store.delete_template(imported["template_id"])
                self.assertIsNotNone(store.get_document(saved["id"]))
                self.assertEqual(store.get_document(saved["id"])["template_name"], source_name.removesuffix(".docx"))
                self.assertFalse(any(row["id"] == imported["template_id"] for row in store.list_templates()))

    def test_subject_suffix_and_summary_use_thai_distributed_layout(self):
        with tempfile.TemporaryDirectory() as temporary:
            target = Path(temporary) / "joined-subject.docx"
            data = document([base_item("market")])
            data["subject"] = "ราคากลางจัดซื้อพัสดุ วัสดุสำนักงาน"
            create_docx(data, target)
            paragraphs = Document(target).paragraphs
            item = next(p for p in paragraphs if "รายการที่ ๑" in p.text)
            summary = next(p for p in paragraphs if p.text.startswith("๓.  คณก.ฯ พิจารณา"))
            self.assertEqual(item.alignment, WD_ALIGN_PARAGRAPH.THAI_JUSTIFY)
            self.assertEqual(summary.alignment, WD_ALIGN_PARAGRAPH.THAI_JUSTIFY)
            self.assertNotIn("\n", item.text)
            self.assertIn("ราคากลางจัดซื้อพัสดุ วัสดุสำนักงาน", "\n".join(p.text for p in paragraphs[:4]))
            self.assertIn("เห็นสมควรใช้ราคาต่ำสุดที่ได้จากการสืบราคา ฯ ในการจัดซื้อพัสดุ วัสดุสำนักงาน รวมเป็นเงินทั้งสิ้น ๕๔๐ บาท (ห้าร้อยสี่สิบบาทถ้วน)", summary.text)
            self.assertNotIn("ราคากลางจัดซื้อพัสดุ วัสดุสำนักงาน จำนวน", summary.text)
            data["items"][0]["quantity"] = "1"
            data["items"][0]["quotes"] = [{"company": "ร้านทดสอบ", "unit_price": "67000"}]
            create_docx(data, target)
            summary = next(p for p in Document(target).paragraphs if p.text.startswith("๓.  คณก.ฯ พิจารณา"))
            self.assertIn("รวมเป็นเงินทั้งสิ้น ๖๗,๐๐๐ บาท (หกหมื่นเจ็ดพันบาทถ้วน)", summary.text)
            self.assertNotIn("จำนวน ๑ รายการ", summary.text)


if __name__ == "__main__":
    unittest.main()
